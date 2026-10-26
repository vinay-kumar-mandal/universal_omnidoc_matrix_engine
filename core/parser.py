# core/parser.py
import os
from pypdf import PdfReader
from docx import Document
from pptx import Presentation
import openpyxl

def split_into_sentences(text):
   
    if not text:
        return []
    
    raw_sentences = text.replace('!', '.').replace('?', '.').split('. ')
    return [s.strip() for s in raw_sentences if len(s.strip()) > 15]

def stream_pdf(file_path):
    reader = PdfReader(file_path)
    for page_num, page in enumerate(reader.pages):
        text = page.extract_text()
        if text:
          
            for sentence in split_into_sentences(text):
                yield (sentence, f"Page {page_num + 1}")

def stream_docx(file_path):
    doc = Document(file_path)
    for idx, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if text:
            for sentence in split_into_sentences(text):
                yield (sentence, f"Paragraph {idx + 1}")

def stream_pptx(file_path):
    prs = Presentation(file_path)
    for slide_idx, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                text = shape.text.strip()
                for sentence in split_into_sentences(text):
                    yield (sentence, f"Slide {slide_idx + 1}")

def stream_excel(file_path):
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        for row_idx, row in enumerate(sheet.iter_rows(values_only=True)):
            row_text = " ".join([str(cell) for cell in row if cell is not None]).strip()
            if row_text:
                
                for sentence in split_into_sentences(row_text):
                    yield (sentence, f"Sheet: {sheet_name}, Row {row_idx + 1}")

def stream_any_document(file_path):
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")
        
    ext = file_path.split('.')[-1].lower()
    
    if ext == 'pdf':
        yield from stream_pdf(file_path)
    elif ext in ['docx', 'doc']:
        yield from stream_docx(file_path)
    elif ext in ['pptx', 'ppt']:
        yield from stream_pptx(file_path)
    elif ext in ['xlsx', 'xls']:
        yield from stream_excel(file_path)
    else:
        raise ValueError(f"Unsupported file format: .{ext}")
